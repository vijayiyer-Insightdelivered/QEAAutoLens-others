import json
import logging
import sqlite3
from datetime import datetime

import openpyxl
import pandas as pd

import config

log = logging.getLogger(__name__)

SCHEMA = """
CREATE TABLE IF NOT EXISTS dealers (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    company_name TEXT NOT NULL,
    company_number TEXT UNIQUE NOT NULL,
    registered_address TEXT,
    postcode TEXT,
    distance_km REAL,
    distance_miles REAL,
    incorporated TEXT,
    companies_house_url TEXT,
    website_url TEXT,
    emails_found TEXT DEFAULT '[]',
    email_source TEXT DEFAULT '{}',
    officers TEXT DEFAULT '[]',
    company_status TEXT,
    status TEXT DEFAULT 'pending',
    error_message TEXT,
    updated_at TEXT
);
"""


def init_db(db_path=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.execute(SCHEMA)
    conn.commit()
    conn.close()
    log.info(f"Database initialized: {db_path}")


def import_from_excel(excel_path, db_path=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    df = pd.read_excel(excel_path)
    conn = sqlite3.connect(db_path)
    cursor = conn.cursor()

    imported = 0
    for _, row in df.iterrows():
        try:
            inc = row.get('Incorporated', '')
            if hasattr(inc, 'strftime'):
                inc = inc.strftime('%Y-%m-%d')
            else:
                inc = str(inc) if pd.notna(inc) else ''

            cursor.execute("""
                INSERT OR IGNORE INTO dealers
                (company_name, company_number, registered_address, postcode,
                 distance_km, distance_miles, incorporated, companies_house_url)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                str(row.get('Company Name', '')),
                str(row.get('Company Number', '')),
                str(row.get('Registered Address', '')),
                str(row.get('Postcode', '')),
                float(row['Distance (km)']) if pd.notna(row.get('Distance (km)')) else None,
                float(row['Distance (miles)']) if pd.notna(row.get('Distance (miles)')) else None,
                inc,
                str(row.get('Companies House URL', '')),
            ))
            imported += cursor.rowcount
        except Exception as e:
            log.warning(f"Failed to import row: {e}")

    conn.commit()
    conn.close()
    total = len(df)
    log.info(f"Imported {imported} new dealers from {total} rows in Excel")
    return imported


def get_pending_batch(db_path=None, batch_size=100):
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    cursor = conn.execute(
        "SELECT * FROM dealers WHERE status = 'pending' ORDER BY id LIMIT ?",
        (batch_size,)
    )
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return rows


def get_dealers_by_numbers(db_path=None, company_numbers=None):
    """Get specific dealers by their company numbers, regardless of status."""
    db_path = db_path or config.DEFAULT_DB_PATH
    if not company_numbers:
        return []
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    placeholders = ','.join('?' for _ in company_numbers)
    cursor = conn.execute(
        f"SELECT * FROM dealers WHERE company_number IN ({placeholders}) ORDER BY id",
        list(company_numbers),
    )
    rows = [dict(r) for r in cursor.fetchall()]
    conn.close()
    return rows


def reset_dealers_by_numbers(db_path=None, company_numbers=None):
    """Reset specific dealers back to pending so they can be re-scraped."""
    db_path = db_path or config.DEFAULT_DB_PATH
    if not company_numbers:
        return 0
    conn = sqlite3.connect(db_path)
    placeholders = ','.join('?' for _ in company_numbers)
    cursor = conn.execute(
        f"UPDATE dealers SET status = 'pending', error_message = NULL WHERE company_number IN ({placeholders})",
        list(company_numbers),
    )
    affected = cursor.rowcount
    conn.commit()
    conn.close()
    return affected


def update_dealer(db_path, company_number, **fields):
    db_path = db_path or config.DEFAULT_DB_PATH
    fields['updated_at'] = datetime.now().isoformat()

    for key in ('emails_found', 'officers'):
        if key in fields and isinstance(fields[key], (list, dict)):
            fields[key] = json.dumps(fields[key])
    if 'email_source' in fields and isinstance(fields['email_source'], dict):
        fields['email_source'] = json.dumps(fields['email_source'])

    set_clause = ', '.join(f'{k} = ?' for k in fields)
    values = list(fields.values()) + [company_number]

    conn = sqlite3.connect(db_path)
    conn.execute(f"UPDATE dealers SET {set_clause} WHERE company_number = ?", values)
    conn.commit()
    conn.close()


def get_progress_stats(db_path=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    cursor = conn.execute("SELECT status, COUNT(*) FROM dealers GROUP BY status")
    stats = dict(cursor.fetchall())
    cursor2 = conn.execute("SELECT COUNT(*) FROM dealers WHERE emails_found != '[]' AND emails_found IS NOT NULL")
    stats['with_emails'] = cursor2.fetchone()[0]
    cursor3 = conn.execute("SELECT COUNT(*) FROM dealers WHERE website_url IS NOT NULL AND website_url != ''")
    stats['with_website'] = cursor3.fetchone()[0]
    conn.close()
    return stats


def export_to_excel(db_path=None, output_path=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    output_path = output_path or config.DEFAULT_OUTPUT_PATH

    conn = sqlite3.connect(db_path)
    df = pd.read_sql_query("SELECT * FROM dealers ORDER BY id", conn)
    conn.close()

    # Format JSON columns for readability
    for col in ('emails_found', 'officers'):
        if col in df.columns:
            df[col] = df[col].apply(_format_json_list)

    if 'email_source' in df.columns:
        df['email_source'] = df['email_source'].apply(_format_json_dict)

    # Drop internal columns
    df.drop(columns=['id', 'updated_at'], inplace=True, errors='ignore')

    df.to_excel(output_path, index=False, engine='openpyxl')
    log.info(f"Exported {len(df)} dealers to {output_path}")
    return output_path


def insert_dealer(db_path=None, **fields):
    """Insert a single dealer from manual entry."""
    db_path = db_path or config.DEFAULT_DB_PATH
    if not fields.get('company_name'):
        raise ValueError("company_name is required")

    # Generate a placeholder company number if not provided
    if not fields.get('company_number'):
        fields['company_number'] = f"MANUAL-{int(datetime.now().timestamp() * 1000)}"

    conn = sqlite3.connect(db_path)
    try:
        conn.execute("""
            INSERT OR IGNORE INTO dealers
            (company_name, company_number, registered_address, postcode,
             companies_house_url, status)
            VALUES (?, ?, ?, ?, ?, 'pending')
        """, (
            fields.get('company_name', ''),
            fields['company_number'],
            fields.get('registered_address', ''),
            fields.get('postcode', ''),
            fields.get('companies_house_url', ''),
        ))
        inserted = conn.total_changes
        conn.commit()
    finally:
        conn.close()
    return fields['company_number']


def get_dealers_paginated(db_path=None, page=1, per_page=50, status_filter=None, search=None):
    """Get paginated dealer list for the dashboard."""
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row

    where_clauses = []
    params = []
    if status_filter and status_filter != 'all':
        where_clauses.append("status = ?")
        params.append(status_filter)
    if search:
        where_clauses.append("(company_name LIKE ? OR company_number LIKE ? OR emails_found LIKE ?)")
        params.extend([f'%{search}%', f'%{search}%', f'%{search}%'])

    where = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

    # Count
    count_row = conn.execute(f"SELECT COUNT(*) FROM dealers {where}", params).fetchone()
    total = count_row[0]

    # Fetch page
    offset = (page - 1) * per_page
    rows = conn.execute(
        f"SELECT * FROM dealers {where} ORDER BY id DESC LIMIT ? OFFSET ?",
        params + [per_page, offset]
    ).fetchall()

    dealers = []
    for r in rows:
        d = dict(r)
        # Parse JSON fields for API response
        for key in ('emails_found', 'officers'):
            if d.get(key):
                try:
                    d[key] = json.loads(d[key])
                except (json.JSONDecodeError, TypeError):
                    pass
        if d.get('email_source'):
            try:
                d['email_source'] = json.loads(d['email_source'])
            except (json.JSONDecodeError, TypeError):
                pass
        dealers.append(d)

    conn.close()
    return {'dealers': dealers, 'total': total, 'page': page, 'per_page': per_page}


def delete_dealer(db_path=None, company_number=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.execute("DELETE FROM dealers WHERE company_number = ?", (company_number,))
    conn.commit()
    conn.close()


def reset_pending(db_path=None):
    """Reset error/processing dealers back to pending."""
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.execute("UPDATE dealers SET status = 'pending' WHERE status IN ('error', 'processing')")
    conn.commit()
    conn.close()


def clear_all(db_path=None):
    db_path = db_path or config.DEFAULT_DB_PATH
    conn = sqlite3.connect(db_path)
    conn.execute("DELETE FROM dealers")
    conn.commit()
    conn.close()


def _format_json_list(val):
    if not val or val == '[]':
        return ''
    try:
        items = json.loads(val)
        return '; '.join(str(x) for x in items) if items else ''
    except (json.JSONDecodeError, TypeError):
        return str(val)


def _format_json_dict(val):
    if not val or val == '{}':
        return ''
    try:
        d = json.loads(val)
        return '; '.join(f'{k}: {v}' for k, v in d.items()) if d else ''
    except (json.JSONDecodeError, TypeError):
        return str(val)
